"""
DOCX Document Parser using python-docx.
"""
from pathlib import Path
from typing import Optional, List
import logging
from docx import Document

from ..base import BaseParser, FileType, DocumentMetadata, ParsedDocument, DocumentSection
from ..normalizer import ContentNormalizer

logger = logging.getLogger(__name__)


class DOCXParser(BaseParser):
    """Extracts paragraphs, headings, and tables from DOCX files."""

    @property
    def supported_type(self) -> FileType:
        return FileType.DOCX

    def parse(self, file_path: Path, metadata: Optional[DocumentMetadata] = None) -> ParsedDocument:
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        file_stat = file_path.stat()
        file_hash = self.calculate_file_hash(file_path)

        if metadata is None:
            metadata = DocumentMetadata(
                filename=file_path.name,
                original_name=file_path.name,
                file_type=FileType.DOCX,
                file_size_bytes=file_stat.st_size,
                file_hash=file_hash,
            )
        else:
            metadata.file_size_bytes = file_stat.st_size
            metadata.file_hash = file_hash

        sections: List[DocumentSection] = []
        raw_sections: List[str] = []
        normalized_sections: List[str] = []

        try:
            doc = Document(str(file_path))
            
            # Extract document core properties if available
            core_props = doc.core_properties
            custom_meta = {}
            if core_props.title:
                custom_meta["title"] = core_props.title
            if core_props.author:
                custom_meta["author"] = core_props.author
            if core_props.comments:
                custom_meta["comments"] = core_props.comments
            metadata.custom_metadata.update(custom_meta)

            section_idx = 0
            current_heading = "Overview"

            # 1. Process paragraphs & headings
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                # Check if paragraph is a heading
                if p.style.name.startswith("Heading"):
                    current_heading = text

                norm_text = ContentNormalizer.normalize(text)
                raw_sections.append(text)
                normalized_sections.append(norm_text)

                sections.append(
                    DocumentSection(
                        section_index=section_idx,
                        section_title=current_heading,
                        raw_text=text,
                        normalized_text=norm_text
                    )
                )
                section_idx += 1

            # 2. Process tables
            for table_idx, table in enumerate(doc.tables):
                table_rows_text: List[str] = []
                for row in table.rows:
                    row_cells = [ContentNormalizer.clean_table_text(cell.text) for cell in row.cells]
                    # Filter out duplicate cells from merged columns
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    if any(unique_cells):
                        table_rows_text.append(" | ".join(unique_cells))

                if table_rows_text:
                    table_content = "\n".join(table_rows_text)
                    raw_sections.append(table_content)
                    normalized_sections.append(table_content)

                    sections.append(
                        DocumentSection(
                            section_index=section_idx,
                            section_title=f"Table {table_idx + 1}",
                            raw_text=table_content,
                            normalized_text=table_content
                        )
                    )
                    section_idx += 1

            full_raw_text = "\n\n".join(raw_sections)
            full_norm_text = "\n\n".join(normalized_sections)

            metadata.char_count = len(full_norm_text)
            metadata.word_count = len(full_norm_text.split()) if full_norm_text else 0
            metadata.status = "parsed"

            return ParsedDocument(
                metadata=metadata,
                raw_text=full_raw_text,
                normalized_text=full_norm_text,
                sections=sections
            )

        except Exception as e:
            metadata.status = "failed"
            metadata.error_message = f"Failed to parse DOCX: {str(e)}"
            logger.error(f"Failed to parse DOCX {file_path}: {e}", exc_info=True)
            raise ValueError(f"Error parsing DOCX file {file_path.name}: {str(e)}") from e
