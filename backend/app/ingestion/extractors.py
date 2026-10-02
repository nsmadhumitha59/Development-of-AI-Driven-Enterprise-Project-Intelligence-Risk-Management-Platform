import io
import os
import pandas as pd
from typing import Tuple, Dict, Any
from pypdf import PdfReader
import docx

def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from a PDF file."""
    reader = PdfReader(io.BytesIO(file_bytes))
    num_pages = len(reader.pages)
    extracted_pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            extracted_pages.append(f"--- Page {i + 1} ---\n{text.strip()}")
    
    full_text = "\n\n".join(extracted_pages)
    extra_metadata = {"num_pages": num_pages}
    return full_text, extra_metadata

def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from a DOCX file."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # Also extract content from tables in DOCX
    table_texts = []
    for t_idx, table in enumerate(doc.tables):
        rows_content = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            rows_content.append(" | ".join(row_cells))
        if rows_content:
            table_texts.append(f"--- Table {t_idx + 1} ---\n" + "\n".join(rows_content))
            
    full_text = "\n\n".join(paragraphs + table_texts)
    extra_metadata = {"paragraph_count": len(paragraphs), "table_count": len(doc.tables)}
    return full_text, extra_metadata

def extract_text_from_csv(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from a CSV file by converting tabular rows into structured semantic sentences."""
    # Attempt parsing CSV with pandas
    df = pd.read_csv(io.BytesIO(file_bytes))
    num_rows, num_cols = df.shape
    columns = list(df.columns)
    
    lines = [f"CSV Document Summary: {num_rows} rows, {num_cols} columns ({', '.join(columns)})."]
    
    # Convert each row into key-value narrative for optimal embedding generation
    for idx, row in df.iterrows():
        row_str_parts = []
        for col in columns:
            val = row[col]
            if pd.notna(val):
                row_str_parts.append(f"{col}: {val}")
        lines.append(f"Record {idx + 1}: " + ", ".join(row_str_parts))
        
    full_text = "\n".join(lines)
    extra_metadata = {"row_count": num_rows, "col_count": num_cols, "columns": columns}
    return full_text, extra_metadata

def extract_text_from_txt(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extract text from a plain TXT file with fallback encoding handling."""
    try:
        text = file_bytes.decode('utf-8')
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode('latin-1')
        except UnicodeDecodeError:
            text = file_bytes.decode('utf-8', errors='ignore')
            
    line_count = len(text.splitlines())
    extra_metadata = {"line_count": line_count}
    return text, extra_metadata

def extract_content(file_bytes: bytes, filename: str) -> Tuple[str, str, Dict[str, Any]]:
    """
    Main extraction dispatcher based on file extension.
    Returns (extracted_text, file_type, extra_metadata)
    """
    ext = os.path.splitext(filename)[1].lower().replace('.', '')
    
    if ext == 'pdf':
        text, meta = extract_text_from_pdf(file_bytes)
        file_type = 'pdf'
    elif ext in ['docx', 'doc']:
        text, meta = extract_text_from_docx(file_bytes)
        file_type = 'docx'
    elif ext == 'csv':
        text, meta = extract_text_from_csv(file_bytes)
        file_type = 'csv'
    elif ext in ['txt', 'md', 'log']:
        text, meta = extract_text_from_txt(file_bytes)
        file_type = 'txt'
    else:
        raise ValueError(f"Unsupported file format '.{ext}'. Supported formats: PDF, DOCX, CSV, TXT.")
        
    if not text or not text.strip():
        raise ValueError(f"File '{filename}' yielded no readable text content.")
        
    return text, file_type, meta
